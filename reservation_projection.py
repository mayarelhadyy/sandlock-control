"""Retryable, non-destructive Excel projection of committed reservations."""
import json
import os
import tempfile
import threading
from contextlib import contextmanager
from pathlib import Path
from openpyxl import load_workbook
from reservation_service import excel_record
from security import safe_text,protect_workbook


@contextmanager
def workbook_lock(path, local_lock):
    with local_lock:
        lock_path=Path(str(path)+'.lock')
        lock_path.parent.mkdir(parents=True,exist_ok=True)
        with open(lock_path,'a+b') as handle:
            handle.seek(0,2)
            if handle.tell()==0:handle.write(b'0');handle.flush()
            handle.seek(0)
            if os.name=='nt':
                import msvcrt
                msvcrt.locking(handle.fileno(),msvcrt.LK_LOCK,1)
            else:
                import fcntl
                fcntl.flock(handle,fcntl.LOCK_EX)
            try:yield
            finally:
                handle.seek(0)
                if os.name=='nt':msvcrt.locking(handle.fileno(),msvcrt.LK_UNLCK,1)
                else:fcntl.flock(handle,fcntl.LOCK_UN)


def atomic_workbook_save(wb,path):
    temp=None
    try:
        fd,temp=tempfile.mkstemp(suffix='.xlsx',dir=Path(path).parent);os.close(fd)
        wb.save(temp);wb.close();os.replace(temp,path);temp=None
    finally:
        wb.close()
        if temp and os.path.exists(temp):os.unlink(temp)


class ReservationProjection:
    def __init__(self, service, path, local_lock=None):
        self.service,self.path=service,Path(path)
        self.local_lock=local_lock or threading.RLock()
        self.error=None

    def flush(self):
        try:pending=[x for x in self.service.store.pending() if not x['projected']]
        except Exception as exc:
            self.error=type(exc).__name__;return False
        if not pending:self.error=None;return True
        try:
            with workbook_lock(self.path,self.local_lock):
                # Fetch after acquiring projection lock: an older projector cannot overwrite newer state.
                records=self.service.list()
                lockers=self.service.availability()
                wb=load_workbook(self.path)
                temp=None
                try:
                    ws=wb['Reservations']
                    headers=[c.value for c in ws[1]]
                    for key in ('User ID','Updated At','Revision','Financial Outcome'):
                        if key not in headers:headers.append(key);ws.cell(1,len(headers)).value=key
                    for r in records:
                        row=excel_record(r)
                        matches=[n for n in range(2,ws.max_row+1) if str(ws.cell(n,1).value or '')==r['bookingId']]
                        if len(matches)>1:raise ValueError('Duplicate legacy Excel rows require migration review')
                        n=matches[0] if matches else ws.max_row+1
                        for i,key in enumerate(headers,1):
                            if key in row:ws.cell(n,i).value=safe_text(row[key])
                    ls=wb['Lockers']
                    for current in lockers:
                        n=next((n for n in range(2,ls.max_row+1) if str(ls.cell(n,1).value)==current['lockerId']),ls.max_row+1)
                        ls.cell(n,1).value=current['lockerId'];ls.cell(n,3).value=current['status'];ls.cell(n,9).value=current['bookingId'] or ''
                    # Preserve the existing Admin cancellation logs as part of the
                    # same atomic workbook replacement, with a retry identity.
                    for event in pending:
                        if event['action']!='cancel' or event['source']!='admin':continue
                        r=json.loads(event['body']);event_id=f"{r['bookingId']}:{r['revision']}"
                        audit=[r['updatedAt'],r.get('actionActor','Owner Dashboard'),'Cancel Reservation',r['bookingId'],r['lockerId'],'committed','Persistent reservation cancelled']
                        access=[r['updatedAt'],r['lockerId'],r['bookingId'],r.get('ownerMobile',''),'Admin Dashboard','cancel-reservation','committed','','admin']
                        for sheet,values in [('Admin Audit',audit),('Access Logs',access)]:
                            log=wb[sheet];keys=[c.value for c in log[1]]
                            if 'Reservation Event' not in keys:
                                keys.append('Reservation Event');log.cell(1,len(keys)).value='Reservation Event'
                            column=keys.index('Reservation Event')+1
                            if any(log.cell(n,column).value==event_id for n in range(2,log.max_row+1)):continue
                            n=log.max_row+1
                            for i,value in enumerate(values,1):log.cell(n,i).value=safe_text(value)
                            log.cell(n,column).value=event_id
                    fd,temp=tempfile.mkstemp(suffix='.xlsx',dir=self.path.parent)
                    os.close(fd);protect_workbook(wb,canonical_ids={r['bookingId'] for r in records});wb.save(temp);wb.close();os.replace(temp,self.path);temp=None
                finally:
                    wb.close()
                    if temp and os.path.exists(temp):os.unlink(temp)
                for item in pending:self.service.store.mark(item['booking_id'],item['revision'],'projected')
            self.error=None
            return True
        except Exception as exc:
            # No rollback of a committed reservation. Durable outbox retries the projection.
            self.error=type(exc).__name__
            return False

    def notify(self, publisher):
        for item in self.service.store.pending():
            if item['notified']:continue
            r=json.loads(item['body'])
            if r['status']!='Cancelled' or publisher(r,item['source']):
                self.service.store.mark(item['booking_id'],item['revision'],'notified')

    def device_audit(self,journal):
        if not journal:return
        events=journal.events()
        if not events:return
        with workbook_lock(self.path,self.local_lock):
            wb=load_workbook(self.path)
            try:
                ws=wb['Admin Audit'];headers=[c.value for c in ws[1]]
                if 'Device Event' not in headers:
                    headers.append('Device Event');ws.cell(1,len(headers)).value='Device Event'
                col=headers.index('Device Event')+1
                existing={str(ws.cell(n,col).value) for n in range(2,ws.max_row+1)}
                for e in events:
                    key=e['request_id']+':'+str(e['event_id'])
                    if key in existing:continue
                    from datetime import datetime,timezone
                    values=[datetime.fromtimestamp(e['timestamp'],timezone.utc).isoformat(),e['actor'],'Device command',e['booking'],e['locker'],e['outcome'],'Physical outcome not certified']
                    n=ws.max_row+1
                    for i,v in enumerate(values,1):ws.cell(n,i).value=safe_text(v)
                    ws.cell(n,col).value=key
                atomic_workbook_save(wb,self.path)
            finally:wb.close()

    def applications(self):
        with workbook_lock(self.path,self.local_lock):
            with self.service.store.connect() as db:
                pending=[dict(r) for r in db.execute('SELECT rowid,* FROM application_events WHERE projected=0 ORDER BY rowid LIMIT 100')]
            if not pending:return True
            wb=load_workbook(self.path)
            try:
                for e in pending:
                    d=json.loads(e['body']);kind=e['kind'];stamp=e['created_at']
                    if kind=='profile':
                        with self.service.store.connect() as db:account=db.execute('SELECT * FROM accounts WHERE user_id=?',(e['actor'],)).fetchone()
                        if not account:continue
                        ws=wb['Users'];n=next((n for n in range(2,ws.max_row+1) if ws.cell(n,1).value==e['actor']),ws.max_row+1)
                        values=[e['actor'],account['name'],account['mobile'],d.get('paymentLast4',''),stamp,stamp,'active']
                        for i,v in enumerate(values,1):ws.cell(n,i).value=safe_text(v)
                        continue
                    if kind=='feedback':sheet='Feedback';values=[stamp,d['bookingId'],d.get('ownerMobile',''),d['rating'],d['comment']]
                    elif kind=='access':sheet='Access Logs';values=[stamp,d['locker'],d['bookingId'],d.get('ownerMobile',''),d['method'],d['action'],d['result'],'unknown',d['source']]
                    elif kind=='audit':sheet='Admin Audit';values=[stamp,e['actor'],d['action'],d.get('bookingId',''),d.get('locker',''),d['result'],d.get('details','')]
                    elif kind=='alert-ack':
                        ws=wb['Alerts'];n=d['row']
                        if n<=ws.max_row and [str(ws.cell(n,i).value or '') for i in range(1,6)]==d['identity']:ws.cell(n,6).value=True
                        else:raise ValueError('Alert identity changed; recovery requires review')
                        continue
                    else:raise ValueError('Unknown application event kind')
                    ws=wb[sheet];headers=[c.value for c in ws[1]]
                    if 'Application Event' not in headers:headers.append('Application Event');ws.cell(1,len(headers)).value='Application Event'
                    col=headers.index('Application Event')+1
                    if any(ws.cell(n,col).value==e['event_id'] for n in range(2,ws.max_row+1)):continue
                    n=ws.max_row+1
                    for i,v in enumerate(values,1):ws.cell(n,i).value=safe_text(v)
                    ws.cell(n,col).value=e['event_id']
                atomic_workbook_save(wb,self.path)
                with self.service.store.connect() as db:
                    db.executemany('UPDATE application_events SET projected=1 WHERE event_id=?',[(e['event_id'],) for e in pending])
            finally:wb.close()
        return True

    def financial(self):
        import payment_service as finance
        try:
            with workbook_lock(self.path,self.local_lock):
                pending=[e for e in finance.records(self.service.store) if not e['projected']]
                if not pending:return True
                values={r['Financial Event']:r for r in finance.reporting_rows(self.service.store)}
                wb=load_workbook(self.path)
                try:
                    ws=wb['Payments'];headers=[c.value for c in ws[1]]
                    for key in ('Financial Event','Financial Outcome','Final Demo Value','Actual Collected Revenue'):
                        if key not in headers:
                            headers.append(key);ws.cell(1,len(headers)).value=key
                    col=headers.index('Financial Event')+1
                    for e in pending:
                        matches=[n for n in range(2,ws.max_row+1) if ws.cell(n,col).value==e['reference']]
                        if len(matches)>1:raise ValueError('Financial projection duplicate requires review')
                        n=matches[0] if matches else ws.max_row+1
                        for i,key in enumerate(headers,1):
                            if key in values[e['reference']]:ws.cell(n,i).value=safe_text(values[e['reference']][key])
                    atomic_workbook_save(wb,self.path)
                    with self.service.store.connect() as db:
                        db.executemany('UPDATE financial_events SET projected=1 WHERE reference=? AND projection_version=?',[(e['reference'],e['projection_version']) for e in pending])
                finally:wb.close()
            return True
        except Exception as exc:
            self.error=type(exc).__name__
            return False
